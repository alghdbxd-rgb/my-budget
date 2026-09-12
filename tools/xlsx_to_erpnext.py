#!/usr/bin/env python3
"""تنظيف كشف مصروفات Excel وتحويله إلى ملفات جاهزة لاستيراد ERPNext.

الاستخدام:
    python3 tools/xlsx_to_erpnext.py <ملف.xlsx> <config.json> <مجلد الإخراج>

الأداة عامة: كل ما يخص ملف بعينه (أسماء الشيتات، أرقام الأعمدة، قواعد توحيد
أسماء الأطراف) يوصف في ملف config.json ولا يُخزَّن هنا.

ما تعالجه:
  * صفوف المجاميع في نهاية كل شيت (SUM) — تُستبعد من البيانات وتُقارن بالمحسوب.
  * التواريخ المختلطة: نص "يوم/شهر/سنة" مقابل خلايا تاريخ فسّرها Excel بنظام
    شهر/يوم. تُحل الالتباسات بالترتيب الزمني للصفوف المجاورة، وما يبقى ملتبساً
    يُعلَّم للمراجعة اليدوية بدل تخمينه بصمت.
  * توحيد أسماء الأطراف (المصادر والمقاولين) وفصل الوصف عن الاسم.
  * رصد الازدواج المحتمل بين صفوف السلف وصفوف الصرف المرتبطة بها.
"""

import csv
import datetime
import json
import os
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

TEXT_DATE = re.compile(r"^\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$")
MONTH_LABEL = re.compile(r"(\d{4})\s*/\s*(\d{1,2})")

CONFIDENT = "مؤكد"
RESOLVED = "مُصحَّح"
NEEDS_REVIEW = "يحتاج مراجعة"
MISSING = "بدون تاريخ"


def cell(row, idx):
    if idx is None or idx >= len(row):
        return None
    v = row[idx]
    return v.strip() if isinstance(v, str) else v


def text(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def last_day(year, month):
    if month == 12:
        return datetime.date(year, 12, 31)
    return datetime.date(year, month + 1, 1) - datetime.timedelta(days=1)


def read_date(value):
    """يرجع (التاريخ, البديل, الحالة). البديل غير فارغ فقط عند الالتباس."""
    if isinstance(value, datetime.datetime):
        d = value.date()
        if d.day > 12 or d.day == d.month:
            return d, None, CONFIDENT
        try:
            return d, datetime.date(d.year, d.day, d.month), NEEDS_REVIEW
        except ValueError:
            return d, None, CONFIDENT
    if isinstance(value, str):
        m = TEXT_DATE.match(value)
        if m:
            day, month, year = (int(x) for x in m.groups())
            try:
                return datetime.date(year, month, day), None, CONFIDENT
            except ValueError:
                return None, None, MISSING
        m = MONTH_LABEL.search(value)
        if m:
            year, month = int(m.group(1)), int(m.group(2))
            if 1 <= month <= 12:
                return last_day(year, month), None, CONFIDENT
    return None, None, MISSING


def disambiguate(entries):
    """يختار القراءة التي تتفق مع ترتيب الصفوف الزمني، ويعلّم ما تبقّى."""
    fixed = [e["date"] if e["date_status"] == CONFIDENT else None for e in entries]
    for i, e in enumerate(entries):
        if e["date_status"] != NEEDS_REVIEW:
            continue
        prev = next((d for d in reversed(fixed[:i]) if d), None)
        nxt = next((d for d in fixed[i + 1:] if d), None)

        def fits(d):
            return (prev is None or d >= prev) and (nxt is None or d <= nxt)

        primary, alt = e["date"], e["date_alt"]
        if fits(primary) and not fits(alt):
            e["date_status"] = CONFIDENT
        elif fits(alt) and not fits(primary):
            e["date"], e["date_alt"] = alt, primary
            e["date_status"] = RESOLVED
        # خلاف ذلك: القراءتان محتملتان أو كلتاهما خارج الترتيب — يبقى للمراجعة
    return entries


def load_sheet(wb, spec, rules):
    ws = wb[spec["name"]]
    entries = []
    skipped = []
    total_row = None
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    for idx, row in enumerate(rows, start=1):
        if idx <= spec["header_row"]:
            continue
        if not any(v not in (None, "") for v in row):
            continue
        amount = cell(row, spec["amount_col"])
        others = [
            cell(row, c)
            for c in range(len(row))
            if c != spec["amount_col"] and cell(row, c) not in (None, "")
        ]
        if isinstance(amount, (int, float)) and not others:
            # صف مجموع؛ الأول هو مجموع الجدول، وما بعده صيغ شاردة في أسفل الورقة
            if total_row is None:
                total_row = amount
            continue
        if not isinstance(amount, (int, float)):
            skipped.append((idx, [text(v) for v in row if text(v)]))
            continue

        raw_date = cell(row, spec.get("date_col")) or cell(row, spec.get("month_label_col"))
        d, alt, status = read_date(raw_date)
        party_raw = text(cell(row, spec.get("party_col")))
        entries.append({
            "excel_row": idx,
            "sheet": spec["name"],
            "category": spec["category"],
            "direction": spec["direction"],
            "date": d,
            "date_alt": alt,
            "date_status": status,
            "date_raw": raw_date.strftime("%Y-%m-%d") if isinstance(raw_date, datetime.datetime) else text(raw_date),
            "party_raw": party_raw,
            "party": normalize_party(party_raw, rules),
            "description": text(cell(row, spec.get("desc_col"))),
            "amount": amount,
            "notes": " | ".join(
                t for t in (text(cell(row, c)) for c in spec.get("note_cols", [])) if t
            ),
        })
    return disambiguate(entries), total_row, skipped


def normalize_party(value, rules):
    if not value:
        return ""
    for pattern, canonical in rules:
        if re.search(pattern, value):
            return canonical
    return value


def audit(all_entries, sheet_totals, duplicate_watch):
    findings = []
    for sheet, declared in sheet_totals.items():
        computed = sum(e["amount"] for e in all_entries if e["sheet"] == sheet)
        if declared is not None and round(declared) != round(computed):
            findings.append(
                f"[{sheet}] مجموع الصفوف {computed:,.0f} لا يطابق صف المجموع في الملف {declared:,.0f}"
                f" (فرق {computed - declared:,.0f})"
            )
    for e in all_entries:
        if any(w in e["description"] or w in e["notes"] for w in duplicate_watch):
            findings.append(
                f"[{e['sheet']} صف {e['excel_row']}] {e['description'] or e['party']} "
                f"{e['amount']:,.0f} — مرتبط بسلفة، تحقق من عدم احتسابه مرتين"
            )
    for e in all_entries:
        if e["date_status"] == MISSING:
            findings.append(f"[{e['sheet']} صف {e['excel_row']}] بدون تاريخ صالح — {e['amount']:,.0f}")
    return findings


HEADERS = [
    "الشيت", "صف الاكسل", "القسم", "الاتجاه", "التاريخ", "حالة التاريخ",
    "التاريخ كما في الملف", "القراءة البديلة", "الطرف", "الاسم كما في الملف",
    "البيان", "المبلغ", "ملاحظات",
]


def row_of(e):
    return [
        e["sheet"], e["excel_row"], e["category"],
        "وارد" if e["direction"] == "in" else "صرف",
        e["date"].strftime("%Y-%m-%d") if e["date"] else "",
        e["date_status"], e["date_raw"],
        e["date_alt"].strftime("%Y-%m-%d") if e["date_alt"] else "",
        e["party"], e["party_raw"], e["description"], e["amount"], e["notes"],
    ]


def write_outputs(entries, findings, sheet_totals, out_dir):
    os.makedirs(out_dir, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "البيانات المنظفة"
    ws.sheet_view.rightToLeft = True
    ws.append(HEADERS)
    for e in entries:
        ws.append(row_of(e))

    review = wb.create_sheet("تواريخ للمراجعة")
    review.sheet_view.rightToLeft = True
    review.append(HEADERS)
    for e in entries:
        if e["date_status"] in (NEEDS_REVIEW, MISSING):
            review.append(row_of(e))

    summary = wb.create_sheet("الملخص")
    summary.sheet_view.rightToLeft = True
    summary.append(["القسم", "عدد الصفوف", "المجموع المحسوب", "مجموع الملف"])
    for sheet, declared in sheet_totals.items():
        rows = [e for e in entries if e["sheet"] == sheet]
        summary.append([sheet, len(rows), sum(r["amount"] for r in rows), declared])
    inflow = sum(e["amount"] for e in entries if e["direction"] == "in")
    outflow = sum(e["amount"] for e in entries if e["direction"] == "out")
    summary.append([])
    summary.append(["إجمالي الوارد", "", inflow])
    summary.append(["إجمالي الصرف", "", outflow])
    summary.append(["الرصيد", "", inflow - outflow])

    issues = wb.create_sheet("ملاحظات التدقيق")
    issues.sheet_view.rightToLeft = True
    issues.append(["الملاحظة"])
    for f in findings:
        issues.append([f])

    xlsx_path = os.path.join(out_dir, "cleaned-data.xlsx")
    wb.save(xlsx_path)

    csv_path = os.path.join(out_dir, "cleaned-data.csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(HEADERS)
        for e in entries:
            w.writerow(row_of(e))

    parties = sorted({e["party"] for e in entries if e["party"]})
    party_path = os.path.join(out_dir, "parties.csv")
    with open(party_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["الاسم الموحّد", "الأسماء الأصلية", "عدد الحركات", "المجموع"])
        for p in parties:
            rows = [e for e in entries if e["party"] == p]
            w.writerow([p, " / ".join(sorted({r["party_raw"] for r in rows})),
                        len(rows), sum(r["amount"] for r in rows)])

    return xlsx_path, csv_path, party_path


# ---------------------------------------------------------------- ERPNext


def acct(name, abbr):
    return f"{name} - {abbr}" if abbr else name


def emit_erpnext(entries, cfg, out_dir):
    """يولّد ملفات Data Import: مراكز الكلفة، الحسابات، وقيود اليومية."""
    erp = cfg.get("erpnext")
    if not erp:
        return []
    abbr, company = erp["abbr"], erp["company"]
    cash = acct(erp["cash_account"], abbr)
    project = erp.get("project", "")
    written = []

    # مراكز الكلفة: مركز لكل قسم تحت المركز الرئيسي
    path = os.path.join(out_dir, "erpnext-01-cost-centers.csv")
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["cost_center_name", "parent_cost_center", "company", "is_group"])
        for cc in erp["cost_centers"]:
            w.writerow([cc, acct(erp["parent_cost_center"], abbr), company, 0])
    written.append(path)

    # الحسابات: حساب مصروف/كلفة لكل قسم + حساب جارٍ لكل ممول
    funders = sorted({e["party"] for e in entries if e["direction"] == "in" and e["party"]})
    path = os.path.join(out_dir, "erpnext-02-accounts.csv")
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["account_name", "parent_account", "company", "is_group",
                    "root_type", "account_type", "account_number"])
        spent = sorted({e["category"] for e in entries if e["direction"] == "out"})
        for i, name in enumerate([erp["category_accounts"][c] for c in spent], start=1):
            w.writerow([name, acct(erp["expense_parent"], abbr), company, 0,
                        erp["expense_root_type"], "", erp["expense_number_prefix"] + str(i)])
        for i, f in enumerate(funders, start=1):
            w.writerow([f"جاري {f}", acct(erp["funding_parent"], abbr), company, 0,
                        "Liability", erp.get("funding_account_type", ""),
                        erp["funding_number_prefix"] + str(i)])
    written.append(path)

    # قيود اليومية: سطر مدين وسطر دائن لكل حركة
    header = ["posting_date", "company", "voucher_type", "user_remark",
              "accounts.account", "accounts.debit_in_account_currency",
              "accounts.credit_in_account_currency", "accounts.cost_center",
              "accounts.project", "accounts.user_remark"]
    buckets = {"ready": [], "review-dates": [], "missing-dates": []}
    for e in entries:
        key = ("missing-dates" if e["date_status"] == MISSING
               else "review-dates" if e["date_status"] == NEEDS_REVIEW else "ready")
        buckets[key].append(e)

    for key, rows in buckets.items():
        if not rows:
            continue
        path = os.path.join(out_dir, f"erpnext-03-journal-entries-{key}.csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(header)
            for e in rows:
                date = e["date"].strftime("%Y-%m-%d") if e["date"] else ""
                label = " — ".join(t for t in (e["description"], e["party"], e["notes"]) if t)[:140]
                cc = acct(erp["category_cost_center"][e["category"]], abbr) if e["category"] in erp["category_cost_center"] else ""
                if e["direction"] == "in":
                    debit = (cash, e["amount"], 0, "")
                    credit = (acct(f"جاري {e['party']}", abbr) if e["party"] else cash, 0, e["amount"], "")
                else:
                    target = acct(erp["category_accounts"][e["category"]], abbr)
                    debit = (target, e["amount"], 0, cc)
                    credit = (cash, 0, e["amount"], "")
                for n, (account, dr, cr, center) in enumerate((debit, credit)):
                    w.writerow([date if n == 0 else "", company if n == 0 else "",
                                "Journal Entry" if n == 0 else "", label if n == 0 else "",
                                account, dr or "", cr or "", center, project, label])
        written.append(path)
    return written


def main():
    if len(sys.argv) != 4:
        print(__doc__)
        return 1
    src, cfg_path, out_dir = sys.argv[1:4]
    cfg = json.load(open(cfg_path, encoding="utf-8"))
    wb = openpyxl.load_workbook(src, data_only=True)

    entries, sheet_totals, findings_skipped = [], {}, []
    for spec in cfg["sheets"]:
        rows, total, skipped = load_sheet(wb, spec, cfg["party_rules"])
        entries.extend(rows)
        sheet_totals[spec["name"]] = total
        for idx, content in skipped:
            findings_skipped.append(f"[{spec['name']} صف {idx}] صف بلا مبلغ رقمي — لم يُستورد: {' | '.join(content)[:90]}")

    findings = findings_skipped + audit(entries, sheet_totals, cfg.get("duplicate_watch", []))
    paths = list(write_outputs(entries, findings, sheet_totals, out_dir))
    paths += emit_erpnext(entries, cfg, out_dir)

    print(f"صفوف مقروءة: {len(entries)}")
    for status in (CONFIDENT, RESOLVED, NEEDS_REVIEW, MISSING):
        n = sum(1 for e in entries if e["date_status"] == status)
        if n:
            print(f"  تواريخ {status}: {n}")
    print(f"ملاحظات تدقيق: {len(findings)}")
    for p in paths:
        print("  ->", p)
    return 0


if __name__ == "__main__":
    sys.exit(main())
